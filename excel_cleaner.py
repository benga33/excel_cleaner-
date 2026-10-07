import argparse
import csv
import io
import math
import re
import sys
import warnings
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAX_WIDTH = 60  

CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f​﻿]")
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
INTL_PHONE = re.compile(r"^(\+|00)\d[\d\s().\-/]{5,}$")
PHONE = re.compile(r"^\+?[\d\s().\-/]{6,}$")
DATE = re.compile(
    r"^(\d{1,4}[./-]\d{1,2}[./-]\d{2,4}\.?"
    r"|\d{1,2}[ .-]?[A-Za-z]{3,9}[ .,-]*\d{2,4}"
    r"|[A-Za-z]{3,9} \d{1,2},? \d{4})"
    r"( \d{1,2}:\d{2}(:\d{2})?)?$"
)
NUMBER = re.compile(r"^[+-]?\d[\d.,]*$")
GROUPED = re.compile(r"^\d+$|^\d{1,3}([.,]\d{3})+$")

EMAIL_WORDS = {"email", "mail"}
PHONE_WORDS = {"phone", "tel", "telefon", "telephone", "mobile", "mob", "gsm",
               "cell", "fax", "mobitel"}
KEEP_TEXT = {"id", "zip", "postal", "postcode", "sku", "code", "iban", "jmbg",
             "pib", "barcode", "ean", "sifra", "šifra", "pbr"}
PLACE_WORDS = {"surname", "prezime", "city", "grad", "town", "mjesto", "country",
               "drzava", "država"}
NAME_WORDS = {"name", "ime"}
NAME_OK = NAME_WORDS | {"first", "last", "full", "given", "family", "middle",
                        "customer", "client", "employee", "student", "contact",
                        "owner", "kupac", "klijent", "korisnik", "i", "and"}



def read_text(path):
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1250"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("latin-1")


def load_sheets(path):
    if path.suffix.lower() in (".csv", ".tsv", ".txt"):
        text = read_text(path)
        try:
            dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect))
        width = max((len(r) for r in rows), default=0)
        return {path.stem: pd.DataFrame([r + [""] * (width - len(r)) for r in rows])}
    return pd.read_excel(path, sheet_name=None, header=None, dtype=str,
                         keep_default_na=False)



def clean_text(value):
    if pd.isna(value):
        return None
    text = CONTROL.sub("", str(value)).replace("\xa0", " ")
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line) or None


def tokens(header):
    spaced = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", header)
    return set(re.findall(r"[^\W\d_]+", spaced.lower()))


def find_header(rows):
    counts = [sum(v is not None for v in r) for r in rows[:20]]
    return next(i for i, c in enumerate(counts) if c >= max(counts) / 2)


def unique_headers(header):
    seen, out = {}, []
    for i, h in enumerate(header, 1):
        h = h or f"Column {i}"
        seen[h] = seen.get(h, 0) + 1
        out.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
    return out


def share(pattern, values):
    return sum(bool(pattern.match(v)) for v in values) / len(values)


def fix_email(v):
    s = v.strip(" ,;<>")
    return s.lower() if EMAIL.match(s) else v


def fix_phone(v):
    v = re.sub(r"\.0$", "", v)  
    if not PHONE.match(v):
        return v  
    digits = re.sub(r"\D", "", v)
    if not 6 <= len(digits) <= 15:
        return v
    if v.startswith("+"):
        return "+" + digits
    if digits.startswith("00"):
        return "+" + digits[2:]
    return digits


def fix_case(v):
    if len(v) <= 2 or not (v.isupper() or v.islower()):
        return v
    return v.title()


def is_name_column(toks):
    if toks & KEEP_TEXT:
        return False
    if toks & PLACE_WORDS:
        return True
    return bool(toks & NAME_WORDS) and toks <= NAME_OK


def parse_numbers(values):
    present = [v for v in values if v is not None]
    if not all(NUMBER.match(v) for v in present):
        return None
    if any(re.match(r"^[+-]?0\d", v) or len(re.sub(r"\D", "", v)) > 15 for v in present):
        return None  
    marks = set()
    for v in present:
        body = v.lstrip("+-")
        if "." in body and "," in body:
            marks.add(max(".,", key=body.rfind)) 
        else:
            for m in ".,":
                if body.count(m) == 1 and len(body) - body.rfind(m) - 1 != 3:
                    marks.add(m)
    if len(marks) > 1:
        return None  
    decimal = marks.pop() if marks else None

    out = []
    for v in values:
        if v is None:
            out.append(None)
            continue
        body = v.lstrip("+-")
        whole, _, fraction = body.partition(decimal) if decimal else (body, "", "")
        if not GROUPED.match(whole):
            return None  
        number = float(re.sub(r"[.,]", "", whole) + ("." + fraction if fraction else ""))
        out.append(-number if v.startswith("-") else number)
    if all(n is None or n == int(n) for n in out):
        out = [None if n is None else int(n) for n in out]
    return out


def clean_column(header, values, dayfirst):
    toks = tokens(header)
    present = [v for v in values if v is not None]
    if not present:
        return "empty", values
    keep_text = bool(toks & KEEP_TEXT)

    if share(EMAIL, present) >= 0.8 or (toks & EMAIL_WORDS and share(EMAIL, present) >= 0.5):
        return "email", [v and fix_email(v) for v in values]

    if toks & PHONE_WORDS or share(INTL_PHONE, present) >= 0.8:
        return "phone", [v and fix_phone(v) for v in values]

    if not keep_text and share(DATE, present) >= 0.8:
        text = pd.Series([v.rstrip(".") if v else None for v in values], dtype=object)
        iso = text.str.match(r"^\d{4}[-/.]\d", na=False)
        parsed = pd.to_datetime(text.where(~iso), format="mixed", dayfirst=dayfirst,
                                errors="coerce")
        parsed = parsed.fillna(pd.to_datetime(text.where(iso), format="mixed",
                                              dayfirst=False, errors="coerce"))
        if parsed.notna().sum() >= 0.8 * len(present):
            return "date", [v if pd.isna(ts) else ts.to_pydatetime()
                            for v, ts in zip(values, parsed)]

    if not keep_text:
        numbers = parse_numbers(values)
        if numbers is not None:
            return "number", numbers

    if is_name_column(toks):
        return "name", [v and fix_case(v) for v in values]

    return "text", values


def process(raw, dayfirst):
    rows = [[clean_text(v) for v in r] for r in raw.itertuples(index=False)]
    filled = [r for r in rows if any(v is not None for v in r)]
    if len(filled) < 2:
        return None
    h = find_header(filled)
    header, body = filled[h], filled[h + 1:]
    if not body:
        return None

    keep = [i for i in range(len(header))
            if header[i] is not None or any(r[i] is not None for r in body)]
    names = unique_headers([header[i] for i in keep])

    columns, kinds = {}, []
    for name, i in zip(names, keep):
        kind, values = clean_column(name, [r[i] for r in body], dayfirst)
        columns[name] = pd.Series(values, dtype=object)  
        kinds.append((name, kind))

    df = pd.DataFrame(columns)
    before = len(df)
    df = df.drop_duplicates().reset_index(drop=True)
    info = {
        "empty_rows": len(rows) - len(filled),
        "skipped": h,
        "duplicates": before - len(df),
        "kinds": kinds,
    }
    return df, info



THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
BAND_FILL = PatternFill("solid", fgColor="F2F5FA")


def style_sheet(ws):
    widths = [8] * ws.max_column
    for r, row in enumerate(ws.iter_rows(), start=1):
        for c, cell in enumerate(row):
            v = cell.value
            horizontal = "left"
            if r == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = HEAD_FILL
                size = len(str(v)) + 4 
            else:
                if isinstance(v, str) and v.startswith("="):
                    cell.data_type = "s"  
                if isinstance(v, datetime):
                    midnight = v.time() == datetime.min.time()
                    cell.number_format = "yyyy-mm-dd" if midnight else "yyyy-mm-dd hh:mm"
                    size = 10 if midnight else 16
                elif isinstance(v, float):
                    cell.number_format = "#,##0.00"
                    size, horizontal = len(f"{v:,.2f}"), "right"
                elif isinstance(v, int):
                    cell.number_format = "0"  
                    size, horizontal = len(str(v)), "right"
                else:
                    size = max((len(x) for x in str(v or "").splitlines()), default=0)
                if r % 2 == 1:
                    cell.fill = BAND_FILL
            cell.alignment = Alignment(horizontal=horizontal, vertical="center",
                                       wrap_text=True)
            cell.border = BORDER
            widths[c] = max(widths[c], size + 2)

    widths = [min(w, MAX_WIDTH) for w in widths]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    for r, row in enumerate(ws.iter_rows(min_row=2), start=2):
        lines = 1
        for c, cell in enumerate(row):
            if isinstance(cell.value, str):
                per_line = max(widths[c] * 0.9, 1)
                lines = max(lines, sum(max(1, math.ceil(len(x) / per_line))
                                       for x in cell.value.splitlines() or [""]))
        if lines > 1:
            ws.row_dimensions[r].height = 15 * lines

    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions



def main():
    warnings.filterwarnings("ignore", category=UserWarning)
    parser = argparse.ArgumentParser(description="Clean a messy spreadsheet.")
    parser.add_argument("source")
    parser.add_argument("output", nargs="?")
    parser.add_argument("--month-first", action="store_true",
                        help="read 05/03/2026 as May 3 instead of 5 March")
    args = parser.parse_args()

    src = Path(args.source)
    if not src.exists():
        sys.exit(f"File not found: {src}")
    if src.suffix.lower() not in {".xlsx", ".xlsm", ".csv", ".tsv", ".txt"}:
        sys.exit("Unsupported file type. Save it as .xlsx or .csv first.")
    dst = Path(args.output) if args.output else src.with_name(f"{src.stem}_clean.xlsx")
    if dst.resolve() == src.resolve():
        sys.exit("Output would overwrite the original. Pick another name.")

    try:
        sheets = load_sheets(src)
    except Exception as error:
        sys.exit(f"Could not read {src.name}: {error}")

    results = {}
    for name, raw in sheets.items():
        result = process(raw, dayfirst=not args.month_first)
        if result:
            results[name[:31]] = result
    if not results:
        sys.exit("No data found in the file.")

    with pd.ExcelWriter(dst, engine="openpyxl") as writer:
        for name, (df, _) in results.items():
            df.to_excel(writer, sheet_name=name, index=False)
            style_sheet(writer.sheets[name])

    for name, (df, info) in results.items():
        print(f'\nSheet "{name}": {len(df)} rows')
        print(f'  Removed {info["empty_rows"]} empty rows and {info["duplicates"]} duplicates. '
              f'Skipped {info["skipped"]} rows above the header.')
        for column, kind in info["kinds"]:
            print(f"  {column[:32]:<32} {kind}")
    print(f"\nSaved to: {dst}")


if __name__ == "__main__":
    main()
